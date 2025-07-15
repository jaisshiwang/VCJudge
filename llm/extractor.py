from pathlib import Path
import fitz  # PyMuPDF
from pptx import Presentation
from utils.logger import setup_logger

logger = setup_logger()

def extract_text_from_pdf(file_path: str) -> str:
    try:
        doc = fitz.open(file_path)
        text = "\n".join([page.get_text() for page in doc])
        doc.close()
        return text.strip()
    except Exception as e:
        logger.error(f"Failed to extract text from PDF: {e}")
        raise

def extract_text_from_pptx(file_path: str) -> str:
    try:
        prs = Presentation(file_path)
        text = "\n".join([
            shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text")
        ])
        return text.strip()
    except Exception as e:
        logger.error(f"Failed to extract text from PPTX: {e}")
        raise

def extract_text(file_path: str) -> str:
    path = Path(file_path)
    if not path.exists():
        msg = f"File not found: {file_path}"
        logger.error(msg)
        raise FileNotFoundError(msg)

    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in [".ppt", ".pptx"]:
        return extract_text_from_pptx(file_path)
    else:
        msg = f"Unsupported file type: {ext}"
        logger.error(msg)
        raise ValueError(msg)